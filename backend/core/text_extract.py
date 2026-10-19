"""Extraction du texte brut d'un CV PDF ou DOCX."""

from pathlib import Path


class UnsupportedFile(Exception):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages).strip()
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(line for line in lines if line.strip()).strip()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise UnsupportedFile("Format non pris en charge : dépose un PDF ou un DOCX.")
